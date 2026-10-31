"""Прогон проверки по Excel-выгрузке каталога и вывод метрик.

Запуск:
    python check_catalog.py [путь_к_xlsx]

Если путь не указан, берётся выгрузка «ТГ Парфюмерия - Категория Духи (все атрибуты).xlsx»
рядом со скриптом (или любой .xlsx со словом «Парфюмерия» в имени).
"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from perfumery_name_check import check_products_with_metrics

DEFAULT_FILE = "ТГ Парфюмерия - Категория Духи (все атрибуты).xlsx"

NAME = "Полное наименование товара (2478)"
TYPE = "Тип парфюмерии (1034)"
CODE = "Код товара"
DATE = "Дата создания"


def find_default_file(base: Path) -> Path:
    exact = base / DEFAULT_FILE
    if exact.is_file():
        return exact
    candidates = sorted(base.glob("*Парфюмерия*.xlsx"))
    return candidates[0] if candidates else exact


def load_products(path: Path) -> list[dict[str, Any]]:
    """Читает нужные колонки выгрузки. Книга гарантированно закрывается."""
    workbook = load_workbook(path, read_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        pos = {title: i for i, title in enumerate(header)}
        missing = [k for k in (NAME, TYPE, CODE, DATE) if k not in pos]
        if missing:
            raise KeyError(f"В выгрузке нет колонок: {missing}")
        return [{k: r[pos[k]] for k in (NAME, TYPE, CODE, DATE)} for r in rows]
    finally:
        workbook.close()


def _date_key(value: Any) -> tuple[datetime, str]:
    """Ключ сортировки по дате: datetime сравниваются как даты, строки — по ISO-разбору.

    Нераспознанные значения считаются самыми ранними; при равенстве решает строка.
    """
    if isinstance(value, datetime):
        return value, ""
    text = "" if value is None else str(value).strip()
    try:
        return datetime.fromisoformat(text), text
    except ValueError:
        return datetime.min, text


def latest_per_code(products: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Для каждого кода товара оставляет самую позднюю карточку."""
    latest: dict[str, dict[str, Any]] = {}
    for product in products:
        code = str(product[CODE])
        if code not in latest or _date_key(product[DATE]) >= _date_key(latest[code][DATE]):
            latest[code] = product
    return list(latest.values())


def show(title: str, items: list[dict[str, Any]]) -> None:
    m = check_products_with_metrics(items)["metrics"]
    print(f"\n=== {title} ===")
    print("Всего:", m["total_rows"])
    print("Совпадения (ok):", m["ok_count"])
    print("Предупреждения (warning):", m["warning_count"])
    print("Ручная проверка:", m["manual_review_count"])
    print("Не проверено:", m["not_checked_count"])
    print("Доля предупреждений, %:", m["warning_rate_percent"])
    print("Доля автоматических решений (ok + warning), %:", m["auto_decision_rate_percent"])


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description="Проверка выгрузки каталога парфюмерии.")
    parser.add_argument("xlsx", nargs="?", type=Path, default=None, help="путь к xlsx-выгрузке")
    args = parser.parse_args(argv)

    path = args.xlsx or find_default_file(Path(__file__).resolve().parent)
    if not path.is_file():
        parser.error(f"файл выгрузки не найден: {path}")

    products = load_products(path)
    show("ПО СТРОКАМ ВЫГРУЗКИ", products)
    show("ПО РАЗНЫМ ТОВАРАМ (код товара)", latest_per_code(products))


if __name__ == "__main__":
    main()
