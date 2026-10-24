from openpyxl import load_workbook
from perfumery_name_check import check_products_with_metrics

FILE = r"ТГ Парфюмерия - Категория Духи (все атрибуты).xlsx"

NAME = "Полное наименование товара (2478)"
TYPE = "Тип парфюмерии (1034)"
CODE = "Код товара"
DATE = "Дата создания"

rows = load_workbook(FILE, read_only=True).active.iter_rows(values_only=True)
header = next(rows)
pos = {title: i for i, title in enumerate(header)}
products = [{k: r[pos[k]] for k in (NAME, TYPE, CODE, DATE)} for r in rows]

# Для подсчёта по разным товарам оставляем самую позднюю карточку каждого кода
latest = {}
for p in products:
    code = str(p[CODE])
    if code not in latest or str(p[DATE]) >= str(latest[code][DATE]):
        latest[code] = p


def show(title, items):
    m = check_products_with_metrics(items)["metrics"]
    print(f"\n=== {title} ===")
    print("Всего:", m["total_rows"])
    print("Совпадения (ok):", m["ok_count"])
    print("Предупреждения (warning):", m["warning_count"])
    print("Ручная проверка:", m["manual_review_count"])
    print("Не проверено:", m["not_checked_count"])
    print("Доля предупреждений, %:", m["warning_rate_percent"])
    print("Покрытие автопроверкой, %:", m["coverage_rate_percent"])


show("ПО СТРОКАМ ВЫГРУЗКИ", products)
show("ПО РАЗНЫМ ТОВАРАМ (код товара)", list(latest.values()))
